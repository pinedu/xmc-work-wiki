# -*- coding: utf-8 -*-
"""
wage-pay-judge 数据源适配器（单一来源的数据读取层）
=================================================
让判定引擎与月报生成器共用同一份数据读取逻辑。判定核心（两轮法）只有一份。

**数据源政策（2026-10-13 老哥拍板）：外部数据源只保留 qxb（黔薪保平台 MCP）**——
standard/zhongyi/db 等文件源与 HTTP 接口源已全部移除；skill 的 references/ 参考文档不受影响。

提供：
  - 通用工具（find_header_row / find_col / to_num / load_xlsx）
  - DataSource 接口基类
  - make_datasource(source, data, config) 工厂（仅 qxb）
  - QxbMcpDataSource 实现在 datasource_qxb.py（黔薪保平台 MCP，streamable HTTP）
"""
import os
import sys
import openpyxl

# ----------------------------- 通用工具（从旧 judge.py 搬入，保持单一实现） -----------------------------
def find_col(header, candidates, exact=False):
    for c in candidates:
        for i, h in enumerate(header):
            if h is None:
                continue
            hs = str(h).strip()
            if exact:
                if hs == c:
                    return i
            else:
                if c in hs:
                    return i
    return None


def find_header_row(rows, candidates):
    best, best_score = 0, 0
    for idx, r in enumerate(rows[:5]):
        score = sum(1 for c in candidates if any(c in str(v) for v in r if v is not None))
        if score > best_score:
            best, best_score = idx, score
    return best


def to_num(v):
    if v is None or v == "":
        return 0.0
    try:
        return float(v)
    except (ValueError, TypeError):
        return 0.0


def load_xlsx(path, sheet=None):
    # 勿用 read_only=True：依赖文件内声明的维度标记，WPS/Excel 编辑后可能陈旧导致只读前几行。
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sh = sheet if (sheet and sheet in wb.sheetnames) else wb.sheetnames[0]
    ws = wb[sh]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return rows


# ----------------------------- 班组名→工种 默认反推词典 -----------------------------
DEFAULT_TEAM_TO_GONG = [
    ("砌筑", "砌筑工"),
    ("模板", "模板工"),
    ("混凝土", "混凝土工"),
    ("水电", "水电工"),
    ("钢筋", "钢筋工"),
    ("机械", "工程机械司机"),
    ("涂料", "涂料工"),
    ("外架", "架子工"),
    ("防水", "防水工"),
    ("杂工", "杂工"),
    ("监理", "管理人员"),
    ("管理", "管理人员"),
]


def team_to_gong(team, mapping=None):
    if not team:
        return None
    t = str(team)
    mp = mapping or DEFAULT_TEAM_TO_GONG
    for kw, g in mp:
        if kw in t:
            return g
    return None


# ----------------------------- 接口基类 -----------------------------
class DataSource:
    """判定引擎与月报生成器共用的最小数据访问接口。"""

    # ---- 判定引擎需要 ----
    def salary_agg(self, ym):
        """当月工资册按身份证汇总：{key: {id,name,team,pay,rows,short}}。"""
        raise NotImplementedError

    def kaoqin(self, ym):
        """当月考勤：(by_nt, by_name)。by_nt={(姓名,班组):天数}，by_name={姓名:天数}。"""
        raise NotImplementedError

    def roster(self):
        """花名册：(persons, by_id, by_nt)。
        persons 每条含 name/id/status/entry/labor/team/gong/contract；
        by_id[身份证]=rec 或 [rec,...]（冲突布局）；by_nt[姓名|班组]=rec 或 [rec,...]。"""
        raise NotImplementedError

    def history(self, target_ym, window=None):
        """历史日薪：({id: [日薪,...]}, prior_months)。"""
        raise NotImplementedError

    def price_table(self):
        """工种单价区间：{工种:(下限,上限)}。"""
        raise NotImplementedError

    # ---- 月报生成器需要 ----
    def list_salary_months(self):
        raise NotImplementedError

    def list_att_months(self):
        raise NotImplementedError

    def attendance_target(self, target_ym):
        """报告侧用：当月考勤 person 列表 [{name,team,days,...}]。"""
        raise NotImplementedError

    # ---- 月报元数据（方案甲：项目概况/资金/落款/制度/投诉统一走数据源）----
    def project_meta(self):
        """返回月报所需的非用工类元数据 dict（键=占位项标签，值=实际值）；缺项返回 {}。
        qxb 源从 qxb_projects_page（项目主数据）取；缺项经生成器合并内联 attach_data 兜底，
        保证判定/渲染始终单一数据源、不引入第二条落盘配置通道。"""
        return {}

    # ---- 月报明细数据接口（2026-10-12：12 类数据接口化；默认空实现=文件源自动降级为【待填】）----
    def offline_attendance(self, target_ym):
        """04 线下考勤册：[{name, team, offline_days}]；未提供返回 []（不做线上线下对比）。"""
        return []

    def fund_flow(self):
        """05 专户进账流水：[{date, payer, amount}]（金额万元）；未提供返回 []。"""
        return []

    def output_value(self):
        """06 产值数据：[{ym, cur_reported, cum_reported, cum_approved}]（万元）；未提供返回 []。"""
        return []

    def six_indicator_status(self, target_ym):
        """07 六项指标状态：[{name, status(达标/限期整改/严重违规), note}]；未提供返回 []。"""
        return []

    def complaint_ledger(self):
        """08 投诉台账：[{date, port(政府端/项目端/薪莫愁端), people, amount, status, handling}]；未提供返回 []。"""
        return []

    def inspections(self):
        """09 迎检记录：[{date, unit, matter, rectify}]；未提供返回 []。"""
        return []

    def entry_exit(self, target_ym):
        """10 进退场与结算：[{name, team, entry, leave, contract, settle_amount, settle_status}]；未提供返回 []。"""
        return []

    def next_attendance(self, next_ym):
        """11 次月已发生考勤：[{name, team, days}]；未提供返回 []（次月预估退回用当月考勤）。"""
        return []

    def yearly_salary(self):
        """12 全年逐月工资册：[{ym, name, team, gong, days, pay}]；未提供返回 []。"""
        return []


# ----------------------------- 工厂 -----------------------------
def make_datasource(source, data, config=None):
    """按 source 选择实现。**数据源政策（2026-10-13）：外部数据源只保留 qxb（黔薪保平台 MCP）**，
    standard/zhongyi/db/custom 均已移除；传入其他值直接报错。

    config 为 dict（可含 price / team_gong_map / qxb 段 / 顶层 project_id）。
    qxb 连接与 token 默认复用 WorkBuddy mcp.json 的 qxb server，不在第二处落盘。
    """
    config = config or {}
    if source in ("qxb", "mcp"):
        here = os.path.dirname(os.path.abspath(__file__))
        if here not in sys.path:
            sys.path.insert(0, here)
        from datasource_qxb import DEFAULT_QXB_CONFIG, QxbMcpDataSource
        qxb_cfg = dict(DEFAULT_QXB_CONFIG)
        user_qxb = config.get("qxb") or {}
        qxb_cfg.update(user_qxb)
        qxb_cfg["project_id"] = (user_qxb.get("project_id")
                                 or config.get("project_id")
                                 or os.environ.get("QXB_PROJECT_ID") or "")
        qxb_cfg["project_name"] = (user_qxb.get("project_name")
                                  or config.get("project_name") or "")
        qxb_cfg["price"] = user_qxb.get("price") or config.get("price") or ""
        return QxbMcpDataSource(data, qxb_cfg)
    raise ValueError(
        f"未知 source: {source!r}（外部数据源仅支持 qxb：黔薪保平台 MCP）")
